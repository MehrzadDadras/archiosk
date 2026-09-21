"""Exercise real browser forms locally or with a human-issued live verification URL.

Default: isolated testing app, in-memory test identity, no production connection.
--live: ARCHIOSK_VERIFICATION_URL must already exist; never issues credentials.
Writes no credential/session material to its proof output. All mutations use
the EVALUATION_INPUT UI; real project inspection is read-only.
"""
import argparse
import json
import os
from pathlib import Path
import sys
import tempfile
import threading
from urllib.parse import urlparse

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


def verify_report_compression(page, base, app, proof, output):
    """Exercise the ordinary evaluation route and observational report in Chromium."""
    from services import survey_evaluation as evaluation
    from services.case_workspace import CaseWorkspaceStore
    from services.runtime_observation import read
    proof['compressed_reports'] = []
    for case in ('matching:fit', 'matching:mandatory-failure', 'matching:partial', 'matching:composition', 'matching:brief'):
        page.goto(base+'/admin/survey-evaluation')
        page.select_option('#case', case)
        page.get_by_role('button', name='Run isolated evaluation', exact=True).click()
        page.wait_for_url('**/attention?analysis=*')
        report = page.locator('#attention-report')
        assert report.is_visible()
        assert page.locator('#investigation-controls').get_attribute('open') is None
        assert 'Governed factual state: UNRESOLVED' in report.inner_text()
        assert report.locator('pre:visible').count() == 0
        entry = page.url
        before = None
        if app:
            location = evaluation.location(app, urlparse(entry).path.split('/')[-2])
            store = CaseWorkspaceStore(location/'registry')
            state_path = store._path_for(evaluation._read(location)['project_id'])
            before = state_path.read_bytes()
        for label in ('Material only', 'Unresolved', 'Conflicts', 'Changes', 'Evidence', 'Technical', 'All'):
            report.get_by_role('link', name=label, exact=True).click()
            selected_url = page.url
            page.get_by_role('link', name='Reload', exact=True).click()
            assert page.url == selected_url
            if before is not None:
                assert state_path.read_bytes() == before
        group = report.locator('[data-report-group]').first
        group.locator(':scope > details > summary').filter(has_text='Show evidence and provenance').click()
        occurrence = group.locator('details').filter(has=page.get_by_text('Occurrence 1', exact=False)).last
        occurrence.locator(':scope > summary').click()
        occurrence.locator('details > summary').click()
        assert occurrence.locator('pre').is_visible()
        source_link = group.get_by_role('link', name='Open source 1', exact=True)
        assert source_link.is_visible()
        source_link.click()
        assert 'Generic kernel mapping' in page.title() or 'Generic kernel mapping' in page.locator('body').inner_text()
        page.goto(entry)
        proof['compressed_reports'].append(dict(case=case, entry=urlparse(entry).path+'?'+urlparse(entry).query,
            default_controls_collapsed=True, filters_and_reload_readonly=True, raw_json_hidden=True,
            complete_occurrence_expanded=True, original_source_jump=True))
    page.screenshot(path=str(output/'compressed-report.png'), full_page=True)
    if app:
        records = [read(app, row['trace']) for row in proof['traces']]
        (output/'runtime-traces.json').write_text(json.dumps(records, indent=2), encoding='utf-8')
        for row, record in zip(proof['traces'], records):
            if row['method'] == 'GET' and row['path'].endswith('/attention'):
                invoked = {event['owner'] for event in record['events'] if event['phase'] == 'INVOKED'}
                assert 'services.case_workspace.CaseWorkspaceStore.project_attention_report' in invoked
                assert 'services.work_product_export.compress_presentation_records' in invoked
                assert not any(owner.endswith(('.run_requirement_matching', '.run_role_composition', '.record_go_attention')) for owner in invoked)
                assert any(event['owner'] == 'components/attention_report.html' and event['phase'] == 'CONSUMED' for event in record['events'])
        proof['report_invoked_consumed_surfaced_without_analysis'] = True
    assert not proof['browser_errors']
    (output/'proof.json').write_text(json.dumps(proof, indent=2), encoding='utf-8')


def verify_work_plans(page, base, app, proof, output):
    """Real two-request UI path: persisted plan, displayed plan, actual executor."""
    from services import survey_evaluation as evaluation
    from services.case_workspace import CaseWorkspaceStore
    from services.runtime_observation import read
    page.select_option('#case', 'matching:fit')
    page.get_by_role('button', name='Run isolated evaluation', exact=True).click()
    page.wait_for_url('**/attention?analysis=*')
    page.locator('#investigation-controls > summary').click()
    form = page.locator('form').filter(has=page.get_by_role('button', name='Set attention', exact=True))
    form.locator('[name=objective]').fill('Inspect source provenance with a declared procedure')
    # Observe the real DOM at the second native submit. This records UI proof
    # only; it does not intercept the request or set application state.
    page.evaluate("""() => document.addEventListener('submit', event => {
      if (event.target.querySelector('[name=plan_id]')) {
        sessionStorage.setItem('work-plan-ui-proof', document.querySelector('#go-work-plans').innerText);
      }
    }, true)""")
    with page.expect_navigation(wait_until='networkidle'):
        form.get_by_role('button', name='Set attention', exact=True).click()
    seen = page.evaluate("sessionStorage.getItem('work-plan-ui-proof')")
    assert seen and 'PLANNED' in seen and 'Procedure only' in seen
    assert page.locator('#go-work-plans > details').count() == 1
    page.locator('#go-work-plans > details > summary').click()
    assert 'PARTIAL' in page.locator('#go-work-plans').inner_text()
    proof['work_plan_displayed_before_execution'] = True
    if app:
        path = evaluation.location(app, urlparse(page.url).path.split('/')[-2])
        store = CaseWorkspaceStore(path/'registry')
        workspace = store.get(evaluation._read(path)['project_id'])
        steps = workspace.investigation_steps
        roots = [s for s in steps if s['step_kind'] == 'governed_work_plan']
        assert len(roots) == 1
        history = [s for s in steps if s.get('branched_from_step_id') == roots[0]['id']]
        assert [s['step_kind'] for s in history] == ['work_plan_execution', 'work_plan_result']
        record = read(app, history[-1]['governed_work_plan']['runtime_trace_id'])
        owner = roots[0]['governed_work_plan']['executor']
        assert any(e['phase'] == 'INVOKED' and e['owner'] == owner for e in record['events'])
        state_path = store._path_for(workspace.project_id)
        before = state_path.read_bytes()
        page.get_by_role('link', name='Reload', exact=True).click()
        assert state_path.read_bytes() == before
        proof['work_plan_persisted_invoked_reload_readonly'] = True
        (output/'runtime-traces.json').write_text(json.dumps([read(app, row['trace']) for row in proof['traces']], indent=2), encoding='utf-8')
    page.screenshot(path=str(output/'work-plan.png'), full_page=True)
    assert not proof['browser_errors']
    (output/'proof.json').write_text(json.dumps(proof, indent=2), encoding='utf-8')


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--live', action='store_true')
    parser.add_argument('--games', action='store_true', help='Exercise continuum and constraint games through real evaluation forms.')
    parser.add_argument('--propositions', action='store_true', help='Exercise sourced subject classification, correction and review through the real UI.')
    parser.add_argument('--matching-games', action='store_true', help='Exercise controlled matching/composition games through the real evaluation UI.')
    parser.add_argument('--report-compression', action='store_true', help='Exercise compact reports, expansion, filters and read-only Reload.')
    parser.add_argument('--work-plans', action='store_true', help='Prove persisted plan display before real runtime execution and read-only Reload.')
    parser.add_argument('--coverage-games', action='store_true', help='Exercise the shared requirement coverage matrix and minimum configurations.')
    parser.add_argument('--transaction-games', action='store_true', help='Exercise retained transaction Claims, real review execution and read-only Reload.')
    parser.add_argument('--cross-domain', action='store_true', help='Compare actual shared-matcher invocation across retained domain cases.')
    parser.add_argument('--output', required=True)
    args=parser.parse_args()
    output=Path(args.output); output.mkdir(parents=True, exist_ok=True)
    server=None; temporary=None
    if args.live:
        access=os.environ.pop('ARCHIOSK_VERIFICATION_URL','')
        parsed=urlparse(access)
        if parsed.scheme!='https' or parsed.hostname not in ('archiosk.com','www.archiosk.com') or not parsed.path.startswith('/verification-access/'):
            raise SystemExit('A human-issued verification URL is required.')
        base=parsed.scheme+'://'+parsed.netloc
    else:
        from app import create_app
        from models import db, User, ROLE_ADMIN
        from werkzeug.security import generate_password_hash
        from werkzeug.serving import make_server
        temporary=tempfile.TemporaryDirectory(prefix='archiosk-source-review-')
        app=create_app('testing')
        app.instance_path=temporary.name
        app.config.update(SECRET_KEY='local-browser-test-only', WTF_CSRF_ENABLED=True,
            REGISTRY_STORE_PATH=str(Path(temporary.name)/'registry'))
        with app.app_context():
            db.create_all()
            user=User(username='local-review-test', role=ROLE_ADMIN)
            user.password_hash=generate_password_hash('Local-Test-Only-2026!')
            db.session.add(user); db.session.commit()
        server=make_server('127.0.0.1',0,app,threaded=True)
        threading.Thread(target=server.serve_forever,daemon=True).start()
        base='http://127.0.0.1:'+str(server.server_port)
    from playwright.sync_api import sync_playwright
    proof=dict(environment='LIVE' if args.live else 'LOCAL_ISOLATED_RUNTIME', traces=[], browser_errors=[])
    try:
        with sync_playwright() as driver:
            browser=driver.chromium.launch(headless=True)
            page=browser.new_page(viewport={'width':1440,'height':1000})
            page.set_default_timeout(120000)
            def reveal_investigation():
                controls = page.locator('#investigation-controls')
                if controls.count() and controls.get_attribute('open') is None:
                    controls.locator(':scope > summary').click()
            def click_and_reveal(locator):
                reveal_investigation()
                planned = locator.evaluate("""element => {
                    const form = element.closest('form'), surface = document.getElementById('go-work-plans');
                    return !!(form && surface && element.type === 'submit' && form.method === 'post' &&
                      JSON.parse(surface.dataset.actions).includes(new FormData(form).get('action') || ''));
                }""")
                if planned:
                    with page.expect_navigation(wait_until='domcontentloaded'):
                        locator.click()
                else:
                    locator.click()
                page.wait_for_load_state('domcontentloaded')
                reveal_investigation()
            page.on('pageerror',lambda error:proof['browser_errors'].append(str(error)))
            def capture(response):
                if response and response.headers.get('x-archiosk-observation'):
                    proof['traces'].append(dict(method=response.request.method,path=urlparse(response.url).path,
                        trace=response.headers['x-archiosk-observation'],status=response.status))
            page.on('response',capture)
            page.goto(base+'/login')
            if args.live:
                page.goto(access); access=''
            else:
                page.fill('#username','local-review-test'); page.fill('#password','Local-Test-Only-2026!')
                click_and_reveal(page.get_by_role('button',name='Sign in',exact=True))
                page.wait_for_load_state('domcontentloaded')
            csrf=page.locator('meta[name="csrf-token"]').get_attribute('content')
            assert page.request.post(base+'/developer-mode/toggle',form={'csrf_token':csrf},headers={'Referer':base+'/'}).ok
            page.goto(base+'/admin/survey-evaluation')
            click_and_reveal(page.get_by_role('button',name='Observe my real requests',exact=True))
            if args.cross_domain:
                identifiers, retained_files = [], []
                for case in ('matching:construction', 'matching:rfp', 'matching:fit', 'matching:asset'):
                    page.goto(base+'/admin/survey-evaluation')
                    page.select_option('#case', case)
                    page.get_by_role('button',name='Run isolated evaluation',exact=True).click()
                    page.wait_for_url('**/attention?analysis=*')
                    identifiers.append(urlparse(page.url).path.split('/')[-2])
                    assert 'Governed factual state: UNRESOLVED' in page.locator('#attention-report').inner_text()
                    if not args.live:
                        from services import survey_evaluation as evaluation
                        from services.case_workspace import CaseWorkspaceStore
                        path = evaluation.location(app, identifiers[-1])
                        record = evaluation._read(path)
                        retained = CaseWorkspaceStore(path/'registry')._path_for(record['project_id'])
                        retained_files.append((retained, retained.read_bytes()))
                page.goto(base+'/admin/survey-evaluation')
                surface = page.locator('#cross-domain-runtime')
                surface.get_by_text('Select retained executions',exact=True).click()
                for identifier in identifiers:
                    surface.locator('input[value="'+identifier+'"]').check()
                surface.get_by_role('button',name='Compare runtime invocations',exact=True).click()
                current_url = page.url
                shared = surface.locator('[data-shared-owner="services.cross_modal_investigation.match_normalized_criteria"]')
                shared.locator(':scope > summary').click()
                for domain in ('construction', 'rfp', 'investment', 'asset'):
                    assert domain+' · INVOKED' in shared.inner_text()
                page.screenshot(path=str(output/'cross-domain-runtime.png'), full_page=True)
                page.get_by_role('button',name='Reload view',exact=True).click()
                assert page.url == current_url
                assert all(path.read_bytes() == before for path, before in retained_files)
                assert not proof['browser_errors']
                proof['cross_domain_actual_shared_invocation'] = True
                proof['reload_preserves_selection_without_analysis'] = True
                proof['entry_path'] = urlparse(current_url).path+'?'+urlparse(current_url).query
                if not args.live:
                    from services.runtime_observation import read
                    (output/'runtime-traces.json').write_text(json.dumps([read(app, row['trace']) for row in proof['traces']],indent=2),encoding='utf-8')
                (output/'proof.json').write_text(json.dumps(proof,indent=2),encoding='utf-8')
                browser.close()
                print(json.dumps(proof,indent=2))
                return
            if args.report_compression:
                verify_report_compression(page, base, None if args.live else app, proof, output)
                browser.close()
                print(json.dumps(proof, indent=2))
                return
            if args.work_plans:
                verify_work_plans(page, base, None if args.live else app, proof, output)
                browser.close()
                print(json.dumps(proof, indent=2))
                return
            if args.transaction_games:
                from services import survey_evaluation as evaluation
                from services.case_workspace import CaseWorkspaceStore
                from services.runtime_observation import read
                proof['transaction_games'] = []
                for name in ('unverified-signing','debt-not-jv','out-of-order'):
                    page.goto(base+'/admin/survey-evaluation')
                    page.select_option('#case','transaction:'+name)
                    page.get_by_role('button',name='Run isolated evaluation',exact=True).click()
                    page.wait_for_url('**/attention?analysis=*')
                    report = page.locator('#attention-report')
                    assert report.is_visible() and report.locator('pre:visible').count() == 0
                    assert 'Governed factual status: UNRESOLVED' in report.inner_text()
                    assert 'Lifecycle: UNRESOLVED' in report.inner_text()
                    reveal_investigation()
                    if name == 'unverified-signing':
                        interpretations = [json.loads(value) for value in page.locator('#transaction-review pre').all_text_contents()]
                        root_claim = next(value for value in interpretations if
                            (value.get('event_proposition') or {}).get('data',{}).get('kind') == 'TRANSACTION_IDENTITY')
                        source = root_claim['event_proposition']
                        data = source['data']
                        source_form = page.locator('#transaction-review form').filter(has=page.locator('input[value="transaction_proposition"]'))
                        source_form.locator('..').locator(':scope > summary').click()
                        for key,value in data['identity']['project'].items():
                            source_form.locator('[name="project_'+key+'"]').fill(value)
                        for key in ('transaction_class','joint_structure'):
                            source_form.locator('[name="'+key+'"]').select_option(data['identity'][key])
                        for key in ('transaction_reference','economic_scope'):
                            source_form.locator('[name="'+key+'"]').fill(data['identity'][key])
                        source_form.get_by_text('Select parties and their source-stated legal identities',exact=True).click()
                        for party in data['identity']['participants']:
                            source_form.locator('[name="party_id"][value="'+party['participant_id']+'"]').check()
                            for key in ('legal_name','jurisdiction','registration_id','role'):
                                source_form.locator('[name="party_'+party['participant_id']+'_'+key+'"]').fill(party[key])
                            source_form.locator('[name="party_'+party['participant_id']+'_identity_relation"]').select_option(party['identity_relation'])
                        for key in ('occurred_at','effective_from','effective_until','discovered_at'):
                            source_form.locator('[name="'+key+'"]').fill(data[key])
                        source_form.locator('[name="evidence_id"]').select_option(source['evidence_fingerprints'][0]['evidence_item_id'])
                        source_form.locator('[name="original_quote"]').fill(source['original_quote'])
                        source_form.locator('[name="source_class"]').select_option('PROJECT_DOCUMENT')
                        source_form.locator('[name="reason"]').fill('Explicit browser-created evaluation interpretation; no authority is granted.')
                        source_form.locator('[name="attribution"]').select_option('agent_assessment')
                        click_and_reveal(source_form.get_by_role('button',name='Record source interpretation',exact=True))
                        assert 'Explicit browser-created evaluation interpretation' in page.locator('#transaction-review').text_content(), page.locator('.flash-message').all_text_contents()
                        proof['source_interpretation_form_invoked'] = True
                    form = page.locator('#transaction-review form').filter(has=page.locator('input[name="action"][value="transaction_review"]'))
                    form.locator('[name="query_date"]').fill('2024-06-01')
                    form.locator('[name="reason"]').fill('Browser-triggered examination of retained hypothetical transaction events.')
                    click_and_reveal(form.get_by_role('button',name='Review retained transaction history',exact=True))
                    entry = page.url
                    before = None
                    if not args.live:
                        location = evaluation.location(app,urlparse(entry).path.split('/')[-2])
                        store = CaseWorkspaceStore(location/'registry')
                        workspace = store.get(evaluation._read(location)['project_id'])
                        assert all(c['event_proposition']['evaluation_only'] for c in workspace.claims)
                        assert not workspace.reviewer_validations and not workspace.applies
                        assert len([r for r in workspace.analyses if (r.get('governed_result') or {}).get('kind') == 'transaction_review']) == 2
                        before = store._path_for(workspace.project_id).read_bytes()
                    page.get_by_role('link',name='Reload',exact=True).click()
                    assert page.url == entry
                    if before is not None:
                        assert store._path_for(workspace.project_id).read_bytes() == before
                    proof['transaction_games'].append(dict(case=name, entry=urlparse(entry).path+'?'+urlparse(entry).query,
                        real_review_form_invoked=True, reload_readonly=True, authority_remains_unresolved=True))
                if not args.live:
                    records = [read(app,t['trace']) for t in proof['traces']]
                    owners = {e['owner'] for r in records for e in r['events'] if e['phase']=='INVOKED'}
                    for owner in ('services.case_workspace.CaseWorkspaceStore.run_transaction_review',
                                  'services.cross_modal_investigation.resolve_transaction_identity',
                                  'services.cross_modal_investigation.investigate_transaction_history'):
                        assert owner in owners
                    proof['invoked_owners'] = sorted(owners)
                assert not proof['browser_errors']
                page.screenshot(path=str(output/'transaction-review.png'),full_page=True)
                (output/'proof.json').write_text(json.dumps(proof,indent=2),encoding='utf-8')
                browser.close()
                print(json.dumps(proof,indent=2))
                return
            if args.coverage_games:
                proof['coverage_games'] = []
                for name in ('coverage-one', 'coverage-two', 'coverage-three', 'coverage-alternatives', 'additive-capital', 'additive-unresolved'):
                    page.goto(base+'/admin/survey-evaluation')
                    page.select_option('#case', 'matching:'+name)
                    page.get_by_role('button', name='Run isolated evaluation', exact=True).click()
                    page.wait_for_url('**/attention?analysis=*')
                    report = page.locator('#attention-report')
                    assert report.is_visible()
                    assert report.locator('nav').evaluate("element => getComputedStyle(element).display") == 'flex'
                    assert 'Governed factual status: UNRESOLVED' in report.inner_text()
                    assert 'Partnership compatibility: UNRESOLVED' in report.inner_text()
                    assert report.locator('pre:visible').count() == 0
                    assert page.locator('#investigation-controls').get_attribute('open') is None
                    if name != 'additive-unresolved':
                        assert 'Minimum sufficient configuration' in report.inner_text()
                        report.locator('summary').filter(has_text='Why fewer participants fail').first.click()
                    else:
                        assert 'No sufficient configuration established' in report.inner_text()
                        assert 'cannot be combined' in report.inner_text()
                    page.get_by_role('link', name='Reload', exact=True).click()
                    proof['coverage_games'].append(dict(case=name, entry=urlparse(page.url).path+'?'+urlparse(page.url).query))
                if not args.live:
                    from services.runtime_observation import read
                    records = [read(app, item['trace']) for item in proof['traces']]
                    assert sum(any(e['phase']=='INVOKED' and e['owner'].endswith('.evaluate_requirement_coverage') for e in r['events']) for r in records) == 6
                    (output/'runtime-traces.json').write_text(json.dumps(records, indent=2), encoding='utf-8')
                assert not proof['browser_errors']
                page.screenshot(path=str(output/'coverage.png'), full_page=True)
                (output/'proof.json').write_text(json.dumps(proof, indent=2), encoding='utf-8')
                browser.close()
                return
            page.select_option('#case','source-review')
            click_and_reveal(page.get_by_role('button',name='Run isolated evaluation',exact=True))
            page.wait_for_load_state('domcontentloaded')
            proof['evaluation_path']=urlparse(page.url).path
            click_and_reveal(page.get_by_role('link',name='Review source frame, anchored machine readings and unresolved stages'))
            page.wait_for_load_state('domcontentloaded')
            proof['review_path']=urlparse(page.url).path
            page.select_option('select[name="view_action"]', 'MIRROR_HORIZONTAL')
            page.locator('form').filter(has=page.locator('select[name="view_action"]')).locator('input[name="reason"]').fill('EVALUATION_INPUT: inspect a mirrored working representation without changing the source.')
            click_and_reveal(page.get_by_role('button',name='Create working view',exact=True))
            assert page.get_by_alt_text('Derived working view; original pixels transformed').count() == 1
            proof['working_view_created_through_ui'] = True
            click_and_reveal(page.get_by_role('button',name='Select four corners on original'))
            image=page.locator('#review-original')
            box=image.bounding_box()
            for x,y in ((.05,.05),(.95,.08),(.9,.95),(.1,.9)):
                image.click(position={'x':box['width']*x,'y':box['height']*y})
            assert len(json.loads(page.input_value('#corners')))==4
            page.fill('#frame-reason','EVALUATION_INPUT: supplied document control quadrilateral; display aspect only.')
            click_and_reveal(page.get_by_role('button',name='Create qualified rectified view',exact=True))
            page.wait_for_load_state('domcontentloaded')
            assert page.get_by_alt_text('Qualified rectified document preview').count()==1
            options=page.locator('#review-field option').evaluate_all('nodes => nodes.map(n => ({value:n.value,text:n.textContent}))')
            field=next(row['value'] for row in options if 'A-2O3' in row['text'])
            page.select_option('#review-field',field)
            page.fill('#after','Sheet: A-203'); page.fill('#text-reason','Controlled source title block visibly reads A-203.')
            click_and_reveal(page.get_by_role('button',name='Propose correction',exact=True))
            click_and_reveal(page.get_by_role('button',name='Accept reviewed correction',exact=True))
            click_and_reveal(page.get_by_role('button',name='Re-evaluate reviewed source',exact=True))
            page.wait_for_load_state('domcontentloaded')
            assert 'Regulatory frontage / front lot line' in page.locator('.survey-evaluation').inner_text()
            before=page.locator('.survey-evaluation').inner_text()
            if not args.live:
                from services import survey_evaluation as evaluation
                from services.case_workspace import CaseWorkspaceStore
                run=proof['evaluation_path'].rsplit('/',1)[1]
                location=evaluation.location(app,run)
                store=CaseWorkspaceStore(str(location/'registry'))
                state_path=store._path_for(evaluation._read(location)['project_id'])
                persisted_before=state_path.read_bytes()
            page.screenshot(path=str(output/'source-review.png'),full_page=True)
            click_and_reveal(page.get_by_role('button',name='Reload view',exact=True))
            page.wait_for_load_state('domcontentloaded')
            after=page.locator('.survey-evaluation').inner_text()
            assert before==after
            if not args.live:
                assert state_path.read_bytes()==persisted_before
                proof['reload_persisted_bytes_unchanged']=True
            click_and_reveal(page.get_by_role('button',name='Revert correction',exact=True))
            assert 'RE-EVALUATION REQUIRED' in page.locator('.survey-evaluation').inner_text()
            click_and_reveal(page.get_by_role('button',name='Re-evaluate reviewed source',exact=True))
            proof['reload_preserves_visible_state']=True
            proof['accept_explicit_reevaluate_revert']=True
            run=proof['evaluation_path'].rsplit('/',1)[1]
            page.goto(base+f'/admin/survey-evaluation/{run}/attention')
            page.fill('input[name="objective"]', 'Inspect sheet identity and retained source readings')
            page.locator('input[name="included_id"]').first.check()
            page.locator('input[name="included_id"]').nth(1).check()
            click_and_reveal(page.get_by_role('button', name='Set attention', exact=True))
            assert 'Persisted attention: SCOPED' in page.locator('body').inner_text()
            scope_text=page.locator('section').first.inner_text()
            click_and_reveal(page.get_by_role('link', name='Reload', exact=True))
            assert page.locator('section').first.inner_text() == scope_text
            proof['attention_invoked_and_reload_preserved_scope']=True
            choices=page.locator('select[name="to_id"] option').evaluate_all('nodes => nodes.map(n => n.value)')
            page.select_option('select[name="to_id"]', choices[1])
            page.fill('input[name="hypothesis"]', 'Possible related sheet readings')
            page.locator('form').filter(has=page.locator('input[value="temporary_relationship"]')).locator('input[name="reason"]').fill('EVALUATION_INPUT: inspect a bounded hypothesis without promoting source authority.')
            click_and_reveal(page.get_by_role('button', name='Add temporary relationship', exact=True))
            assert 'TEMPORARY · Possible related sheet readings' in page.locator('body').inner_text()
            proof['temporary_relationship_invoked_through_ui']=True
            def action_form(action):
                reveal_investigation()
                return page.locator('form').filter(has=page.locator(f'input[name="action"][value="{action}"]'))
            def open_coverage():
                reveal_investigation()
                detail=page.get_by_text('Record conditions and corresponding representations for coverage review', exact=True).locator('..')
                if detail.get_attribute('open') is None:
                    click_and_reveal(detail.locator('summary'))
            open_coverage()
            form=action_form('record_condition')
            assert form.locator('select[name="evidence_id"] option').count() > 0
            form.locator('input[name="meaning"]').fill('EVALUATION_INPUT: located review condition')
            form.locator('input[name="affected_disciplines"]').fill('architectural, mechanical')
            form.locator('select[name="required_resolution"]').select_option('ASSEMBLY_LAYERS')
            form.locator('input[name="reason"]').fill('Evaluation of explicit anchored review; no claim of automatic condition recognition.')
            click_and_reveal(form.locator('button'))
            open_coverage()
            form=action_form('confirm_condition')
            form.locator('input[name="reason"]').fill('EVALUATION_INPUT: confirm the bounded review condition and scope.')
            click_and_reveal(form.locator('button'))
            open_coverage()
            form=action_form('record_representation')
            form.locator('input[name="discipline"]').fill('architectural')
            form.locator('select[name="representation_class"]').select_option('WALL_SECTION')
            form.locator('select[name="resolution_class"]').select_option('ASSEMBLY_LAYERS')
            form.locator('input[name="reason"]').fill('EVALUATION_INPUT: explicit applicability premise for this bounded condition.')
            click_and_reveal(form.locator('button'))
            open_coverage()
            form=action_form('resolve_representation')
            form.locator('input[name="reason"]').fill('EVALUATION_INPUT: reviewed applicability; source authority remains unchanged.')
            click_and_reveal(form.locator('button'))
            form=action_form('professional_review')
            form.locator('select[name="narrative"]').select_option('building_science')
            form.locator('input[name="subject"]').fill('EVALUATION_INPUT envelope interface')
            form.locator('input[name="project_phase"]').fill('design-review')
            form.locator('input[name="discipline"]').fill('architectural')
            form.locator('select[name="current_resolution"]').select_option('OVERALL')
            form.locator('select[name="participation_expectation"]').select_option('upstream')
            form.locator('select[name="required_resolution"]').select_option('LOCAL_TIE_IN')
            form.locator('input[name="reason"]').fill('An overall representation cannot establish the exact seal termination.')
            click_and_reveal(form.locator('button'))
            assert 'INSUFFICIENT_SCALE' in page.locator('body').inner_text()
            assert 'DISCIPLINE_COVERAGE_GAP' in page.locator('body').inner_text()
            click_and_reveal(action_form('professional_presentation').locator('button'))
            with page.expect_download() as download_info:
                click_and_reveal(page.locator('a[href*="presentations/"][href$=".docx"]:visible'))
            download=download_info.value
            report_path=output/'professional-review.docx'
            download.save_as(str(report_path))
            import docx
            report_text='\n'.join(p.text for p in docx.Document(report_path).paragraphs)
            assert 'INSUFFICIENT_SCALE' in report_text and 'DISCIPLINE_COVERAGE_GAP' in report_text
            proof['retained_presentation_formats'] = []
            for extension, label in [('pdf', 'PDF'), ('pptx', 'PowerPoint'), ('html', 'Web report')]:
                with page.expect_download() as format_download:
                    click_and_reveal(page.get_by_role('link', name='Download ' + label, exact=True).filter(visible=True).first)
                artifact_path = output / ('professional-review.' + extension)
                format_download.value.save_as(str(artifact_path))
                if extension == 'pdf':
                    import fitz
                    with fitz.open(artifact_path) as artifact:
                        rendered = '\n'.join(p.get_text() for p in artifact)
                elif extension == 'pptx':
                    from pptx import Presentation
                    rendered = '\n'.join(shape.text for slide in Presentation(artifact_path).slides
                        for shape in slide.shapes if shape.has_text_frame)
                else:
                    rendered = artifact_path.read_text(encoding='utf-8')
                assert 'INSUFFICIENT_SCALE' in rendered and 'DISCIPLINE_COVERAGE_GAP' in rendered
                assert 'EVALUATION_INPUT' in rendered or 'Evaluation Only: True' in rendered
                proof['retained_presentation_formats'].append(dict(format=extension, real_download=True, unresolved_preserved=True))
            proof['professional_review_coverage_and_retained_report_invoked_through_ui']=True
            form=action_form('constraint_review')
            form.locator('input[name="subject"]').fill('EVALUATION_INPUT opening')
            form.locator('input[name="parameter"]').fill('width')
            form.locator('input[name="unit"]').fill('mm')
            for name, value in {'lower_1':'990','upper_1':'1010','lower_2':'1000','upper_2':'1020','baseline':'1005'}.items():
                form.locator(f'input[name="{name}"]').fill(value)
            form.locator('input[name="reason"]').fill('Explicit hypothetical bounds; no project measurement is inferred.')
            click_and_reveal(form.locator('button'))
            assert 'BOUNDARY_FOUND' in page.locator('body').inner_text()
            assert 'Actual project value: UNRESOLVED' in page.locator('body').inner_text()
            click_and_reveal(page.get_by_role('link', name='Reload', exact=True))
            assert 'BOUNDARY_FOUND' in page.locator('body').inner_text()
            proof['hypothetical_constraint_and_breakpoint_invoked_through_ui']=True
            form=action_form('information_comparison')
            form.locator('input[name="property_key"]').fill('opening-width')
            for side in ('left', 'right'):
                form.locator(f'input[name="{side}_scope"]').fill('evaluation-interface')
                form.locator(f'input[name="{side}_value"]').fill('100')
                form.locator(f'input[name="{side}_unit"]').fill('mm')
                form.locator(f'input[name="{side}_qualifiers"]').fill('[]')
                form.locator(f'select[name="{side}_view"]').select_option('VIEW_INVARIANT')
            mirrored=form.locator('select[name="left_view"] option').evaluate_all(
                'nodes => nodes.filter(n => n.textContent.includes("MIRROR_HORIZONTAL")).map(n => n.value)')
            assert len(mirrored) == 1
            form.locator('select[name="left_view"]').select_option(mirrored[0])
            form.locator('input[name="reason"]').fill('EVALUATION_INPUT: compare explicit width hypotheses using the retained mirrored view; no factual consistency is inferred.')
            click_and_reveal(form.locator('button'))
            assert 'model MATCH' in page.locator('body').inner_text()
            assert 'Factual consistency: UNRESOLVED' in page.locator('body').inner_text()
            if not args.live:
                persisted_before=state_path.read_bytes()
                workspace=store.get(evaluation._read(location)['project_id'])
                comparison=workspace.analyses[-1]['governed_result']
                assert comparison['views'][0]['transform']['type'] == 'MIRROR_HORIZONTAL'
                assert comparison['evaluation_only'] and not comparison['canonical']
            click_and_reveal(page.get_by_role('link', name='Reload', exact=True))
            assert 'model MATCH' in page.locator('body').inner_text()
            if not args.live:
                assert state_path.read_bytes() == persisted_before
            proof['normalized_comparison_retained_view_and_reload_invoked_through_ui']=True
            attention_url=page.url
            click_and_reveal(page.get_by_role('link', name='Inspect muscles for this execution', exact=True).last)
            inspector=page.locator('#muscle-inspector')
            assert inspector.is_visible()
            comparison_contract=inspector.locator('[data-muscle-owner="services.cross_modal_investigation.compare_normalized_information"]')
            assert comparison_contract.get_attribute('data-invocation') == 'INVOKED'
            click_and_reveal(comparison_contract.locator('summary').first)
            assert 'Factual consistency: UNRESOLVED' in comparison_contract.inner_text()
            assert 'UNRESOLVED' in comparison_contract.inner_text()
            click_and_reveal(comparison_contract.get_by_text('Recorded inputs / result / provenance', exact=True).first)
            assert 'opening-width' in comparison_contract.inner_text()
            comparison_contract.scroll_into_view_if_needed()
            page.screenshot(path=str(output/'muscle-inspector.png'))
            inspector_text=inspector.text_content()
            click_and_reveal(page.get_by_role('button', name='Reload view', exact=True))
            assert page.locator('#muscle-inspector').text_content() == inspector_text
            page.goto(attention_url)
            if not args.live:
                assert state_path.read_bytes() == persisted_before
            proof['linked_muscle_inspector_actual_invocation_and_read_only_reload']=True
            if args.propositions:
                click_and_reveal(page.get_by_text('Add a subject reference', exact=True))
                form=action_form('record_subject')
                form.locator('input[name="subject_name"]').fill('EVALUATION subject')
                form.locator('input[name="subject_role"]').fill('institutional_investor')
                click_and_reveal(form.locator('button'))
                click_and_reveal(page.get_by_text('Record a sourced proposition', exact=True))
                form=action_form('subject_proposition')
                evidence_id=form.locator('select[name="evidence_id"]').input_value()
                citation=page.locator('a[href*="item=evidence_items:'+evidence_id+'"]').first.get_attribute('href')
                page.goto(base+citation)
                click_and_reveal(page.locator('details').filter(has=page.locator('#kernel-record')).locator('summary'))
                source_quote=json.loads(page.locator('#kernel-record').inner_text())['content'][:200]
                page.goto(attention_url)
                click_and_reveal(page.get_by_text('Record a sourced proposition', exact=True))
                form=action_form('subject_proposition')
                subject=form.locator('select[name="subject_key"] option').evaluate_all(
                    'nodes => nodes.find(n => n.textContent.includes("EVALUATION subject")).value')
                form.locator('select[name="subject_key"]').select_option(subject)
                form.locator('input[name="property_key"]').fill('declared-review-topic')
                form.locator('input[name="scope_key"]').fill('evaluation-source-classification')
                form.locator('select[name="kind"]').select_option('TOKEN_SET')
                form.locator('input[name="value"]').fill('EVALUATION_TOPIC')
                form.locator('input[name="vocabulary"]').fill('evaluation-topics')
                form.locator('input[name="qualifiers"]').fill('[]')
                form.locator('select[name="view_basis"]').select_option('VIEW_INVARIANT')
                form.locator('select[name="evidence_id"]').select_option(evidence_id)
                form.locator('textarea[name="original_quote"]').fill(source_quote)
                form.locator('select[name="source_class"]').select_option('PROJECT_DOCUMENT')
                form.locator('select[name="temporal_class"]').select_option('HISTORICAL_ACTIVITY')
                form.locator('input[name="as_of"]').fill('2020-01-01')
                form.locator('textarea[name="reason"]').fill('EVALUATION_INPUT categorization only; no investor fact or current mandate is established.')
                form.locator('input[name="attribution"][value="agent_assessment"]').check()
                click_and_reveal(form.locator('button'))
                rows=page.locator('[data-proposition-id]')
                assert rows.count() == 1 and 'EVALUATION_INPUT' in rows.first.inner_text()
                original_id=rows.first.get_attribute('data-proposition-id')
                if not args.live:
                    # A local test identity exercises the human form contract.
                    # Automated live verification must not attest human review.
                    rows.first.locator('input[name="reason"]').fill('Retain this evaluation interpretation only.')
                    rows.first.locator('input[name="attribution"]').check()
                    click_and_reveal(rows.first.get_by_role('button', name='Accept as interpretation', exact=True))
                assert 'NOT_ESTABLISHED' in page.locator('[data-proposition-id]').first.inner_text()
                click_and_reveal(page.get_by_role('link', name='Correct this proposition', exact=True))
                form=action_form('subject_proposition')
                assert form.locator('textarea[name="original_quote"]').input_value() == source_quote
                form.locator('input[name="value"]').fill('EVALUATION_UPDATED_TOPIC')
                form.locator('textarea[name="reason"]').fill('Correct the evaluation categorization without changing its source.')
                form.locator('input[name="attribution"][value="agent_assessment"]').check()
                click_and_reveal(form.get_by_role('button', name='Record corrected successor', exact=True))
                rows=page.locator('[data-proposition-id]')
                assert rows.count() == 2
                assert 'superseded' in page.locator('[data-proposition-id="'+original_id+'"]').inner_text()
                if not args.live:
                    rows.last.locator('input[name="reason"]').fill('Reject the proposed evaluation categorization.')
                    rows.last.locator('input[name="attribution"]').check()
                    click_and_reveal(rows.last.get_by_role('button', name='Reject interpretation', exact=True))
                    assert 'rejected' in page.locator('[data-proposition-id]').last.inner_text()
                if not args.live:
                    persisted_before=state_path.read_bytes()
                click_and_reveal(page.get_by_role('link', name='Reload', exact=True))
                assert page.locator('[data-proposition-id]').count() == 2
                if not args.live:
                    assert state_path.read_bytes() == persisted_before
                page.locator('#subject-propositions').scroll_into_view_if_needed()
                page.screenshot(path=str(output/'subject-propositions.png'),full_page=True)
                proof['subject_proposition_create_correct_reload']=True
                proof['subject_proposition_human_curation_local_fixture_only']=not args.live
                matching_claims=[]
                for temporal in ('DATED_REQUIREMENT', 'CURRENT_DISCLOSED_MANDATE'):
                    click_and_reveal(page.get_by_text('Record a sourced proposition', exact=True))
                    form=action_form('subject_proposition')
                    form.locator('select[name="subject_key"]').select_option(subject)
                    form.locator('input[name="property_key"]').fill('evaluation-sector')
                    form.locator('input[name="scope_key"]').fill('evaluation-mandate')
                    form.locator('select[name="kind"]').select_option('TOKEN_SET')
                    form.locator('input[name="value"]').fill('EVALUATION_SECTOR')
                    form.locator('input[name="vocabulary"]').fill('evaluation-sectors')
                    form.locator('input[name="qualifiers"]').fill('[]')
                    form.locator('select[name="view_basis"]').select_option('VIEW_INVARIANT')
                    form.locator('select[name="evidence_id"]').select_option(evidence_id)
                    form.locator('textarea[name="original_quote"]').fill(source_quote)
                    form.locator('select[name="source_class"]').select_option('PROJECT_DOCUMENT')
                    form.locator('select[name="temporal_class"]').select_option(temporal)
                    form.locator('input[name="as_of"]').fill('2026-01-01')
                    form.locator('input[name="valid_until"]').fill('2027-01-01')
                    form.locator('textarea[name="reason"]').fill('EVALUATION_INPUT typed premise only; not an investor fact.')
                    form.locator('input[name="attribution"][value="agent_assessment"]').check()
                    click_and_reveal(form.locator('button'))
                    matching_claims.append(page.locator('[data-proposition-id]').last.get_attribute('data-proposition-id'))
                form=action_form('requirement_matching')
                form.locator('select[name="context_key"]').select_option('investment')
                form.locator('select[name="target_subject"]').select_option(subject)
                form.locator('input[name="query_date"]').fill('2026-09-20')
                form.locator('select[name="criterion_0_required"]').select_option(matching_claims[0])
                form.locator('select[name="criterion_0_candidate"]').select_option(matching_claims[1])
                form.locator('select[name="criterion_0_operator"]').select_option('CONTAINS_ALL')
                form.locator('input[name="reason"]').fill('Compare declared evaluation premises; preserve unresolved factual fit.')
                click_and_reveal(form.get_by_role('button', name='Run requirement matching', exact=True))
                matching=page.locator('#requirement-matching')
                assert 'conditional model FIT' in matching.inner_text()
                assert 'Governed factual fit: UNRESOLVED' in matching.inner_text()
                if not args.live:
                    persisted_before=state_path.read_bytes()
                click_and_reveal(page.get_by_role('link', name='Reload', exact=True))
                assert 'conditional model FIT' in matching.inner_text()
                if not args.live:
                    assert state_path.read_bytes() == persisted_before
                matching.scroll_into_view_if_needed()
                page.screenshot(path=str(output/'requirement-matching.png'),full_page=True)
                proof['requirement_matching_conditional_fit_and_read_only_reload']=True
                click_and_reveal(page.get_by_text('Record a sourced proposition', exact=True))
                form=action_form('subject_proposition')
                form.locator('select[name="subject_key"]').select_option(subject)
                form.locator('input[name="property_key"]').fill('complete_requirement_inventory')
                form.locator('input[name="scope_key"]').fill('evaluation-source-classification')
                form.locator('select[name="kind"]').select_option('TOKEN_SET')
                form.locator('input[name="value"]').fill(matching_claims[0])
                form.locator('input[name="vocabulary"]').fill('requirement_claims')
                form.locator('input[name="qualifiers"]').fill('[]')
                form.locator('select[name="view_basis"]').select_option('VIEW_INVARIANT')
                form.locator('select[name="evidence_id"]').select_option(evidence_id)
                form.locator('textarea[name="original_quote"]').fill(source_quote)
                form.locator('select[name="source_class"]').select_option('PROJECT_DOCUMENT')
                form.locator('select[name="temporal_class"]').select_option('DATED_REQUIREMENT')
                form.locator('input[name="as_of"]').fill('2026-01-01')
                form.locator('input[name="valid_until"]').fill('2027-01-01')
                form.locator('textarea[name="reason"]').fill('EVALUATION_INPUT inventory proposal; no complete governing scope is established.')
                form.locator('input[name="attribution"][value="agent_assessment"]').check()
                click_and_reveal(form.locator('button'))
                inventory=page.locator('[data-proposition-id]').last.get_attribute('data-proposition-id')
                form=action_form('requirement_matching')
                form.locator('select[name="context_key"]').select_option('investment')
                form.locator('select[name="target_subject"]').select_option(subject)
                form.locator('input[name="query_date"]').fill('2026-09-20')
                form.locator('select[name="criterion_0_required"]').select_option(matching_claims[0])
                form.locator('select[name="criterion_0_candidate"]').select_option(matching_claims[1])
                form.locator('select[name="criterion_0_operator"]').select_option('CONTAINS_ALL')
                click_and_reveal(form.get_by_text('Reviewed factual scope', exact=True))
                form.locator('select[name="inventory_claim_id"]').select_option(inventory)
                form.locator('input[name="reason"]').fill('Exercise scoped admission; hypothetical inventory must remain unresolved.')
                click_and_reveal(form.get_by_role('button', name='Run requirement matching', exact=True))
                assert 'Complete requirement inventory authority or applicability is not established.' in page.locator('body').inner_text()
                assert 'Reviewed inventory result: UNRESOLVED' in matching.inner_text()
                proof['reviewed_scope_form_invoked_without_fabricated_authority']=True
                form=action_form('role_composition')
                form.locator('select[name="coverage_mode"]').select_option('reviewed_requirements')
                form.locator('input[name="matching_id"]').last.check()
                form.locator('input[name="required_role_id"][value="'+matching_claims[0]+'"]').check()
                form.locator('input[name="reason"]').fill('Exercise reviewed composition; evaluation evidence cannot establish factual coverage.')
                click_and_reveal(form.get_by_role('button', name='Compose role coverage', exact=True))
                assert 'CONFIGURATION_UNRESOLVED' in page.locator('#attention-report').inner_text()
                assert 'A common reviewed complete inventory is required.' in page.locator('#attention-report').inner_text()
                proof['reviewed_composition_form_invoked_without_fabricated_authority']=True
                form=action_form('professional_review')
                form.locator('select[name="narrative"]').select_option('building_science')
                form.locator('select[name="focus_id"]').select_option(evidence_id)
                form.locator('input[name="subject"]').fill('Evaluation interpretation chain')
                form.locator('input[name="project_phase"]').fill('review')
                form.locator('input[name="discipline"]').fill('architectural')
                click_and_reveal(form.get_by_text('Interpretation chain and proposition root trace', exact=True))
                form.locator('select[name="upstream_claim_id"]').select_option(matching_claims[0])
                form.locator('select[name="target_claim_id"]').select_option(matching_claims[1])
                form.locator('input[name="query_date"]').fill('2026-09-20')
                form.locator('input[name="reason"]').fill('Retain conditional meaning and unresolved authority for evaluation propositions.')
                click_and_reveal(form.get_by_role('button',name='Run professional review',exact=True))
                assert 'Proposition root / return: UNRESOLVED' in page.locator('#attention-report').inner_text()
                proof['interpretation_drift_and_proposition_root_real_form']=True
                click_and_reveal(page.get_by_text('Add a subject reference', exact=True))
                form=action_form('record_subject')
                form.locator('input[name="subject_name"]').fill('EVALUATION debt participant')
                form.locator('input[name="subject_role"]').fill('lender')
                click_and_reveal(form.locator('button'))
                click_and_reveal(page.get_by_text('Record a sourced proposition', exact=True))
                options=action_form('subject_proposition').locator('select[name="subject_key"] option')
                debtor=options.evaluate_all('nodes => nodes.find(n => n.textContent.includes("EVALUATION debt participant")).value')
                opportunity=options.evaluate_all('nodes => nodes.find(n => n.value.startsWith("source:")).value')
                role_claims=[]
                for role_subject, role_value in [(opportunity,'EQUITY'),(opportunity,'DEBT'),(subject,'EQUITY'),(debtor,'DEBT')]:
                    page.goto(attention_url)
                    click_and_reveal(page.get_by_text('Record a sourced proposition', exact=True))
                    form=action_form('subject_proposition')
                    form.locator('select[name="subject_key"]').select_option(role_subject)
                    form.locator('input[name="property_key"]').fill('role')
                    form.locator('input[name="scope_key"]').fill('evaluation-opportunity')
                    form.locator('select[name="kind"]').select_option('TOKEN_SET')
                    form.locator('input[name="value"]').fill(role_value)
                    form.locator('input[name="vocabulary"]').fill('evaluation-roles')
                    form.locator('input[name="qualifiers"]').fill('[]')
                    form.locator('select[name="view_basis"]').select_option('VIEW_INVARIANT')
                    form.locator('select[name="evidence_id"]').select_option(evidence_id)
                    form.locator('textarea[name="original_quote"]').fill(source_quote)
                    form.locator('select[name="source_class"]').select_option('PROJECT_DOCUMENT')
                    form.locator('select[name="temporal_class"]').select_option('CURRENTNESS_UNRESOLVED')
                    form.locator('textarea[name="reason"]').fill('EVALUATION_INPUT role premise only; not a verified financial commitment.')
                    form.locator('input[name="attribution"][value="agent_assessment"]').check()
                    click_and_reveal(form.locator('button'))
                    role_claims.append(page.locator('[data-proposition-id]').last.get_attribute('data-proposition-id'))
                role_matches=[]
                for target, candidate_claim in [(subject,role_claims[2]),(debtor,role_claims[3])]:
                    form=action_form('requirement_matching')
                    form.locator('select[name="context_key"]').select_option('investment')
                    form.locator('select[name="target_subject"]').select_option(target)
                    form.locator('select[name="require_currentness"]').select_option('no')
                    click_and_reveal(form.locator('details').nth(1).locator('summary'))
                    for index in (0,1):
                        prefix='criterion_'+str(index)+'_'
                        form.locator('select[name="'+prefix+'required"]').select_option(role_claims[index])
                        form.locator('select[name="'+prefix+'candidate"]').select_option(candidate_claim)
                        form.locator('select[name="'+prefix+'operator"]').select_option('CONTAINS_ALL')
                        form.locator('select[name="'+prefix+'mandatory"]').select_option('no')
                    form.locator('input[name="reason"]').fill('Evaluate the role contribution of this candidate.')
                    click_and_reveal(form.get_by_role('button',name='Run requirement matching',exact=True))
                    role_matches.append(page.locator('input[name="matching_id"]').last.input_value())
                form=action_form('role_composition')
                for identifier in role_claims[:2]:
                    form.locator('input[name="required_role_id"][value="'+identifier+'"]').check()
                for identifier in role_matches:
                    form.locator('input[name="matching_id"][value="'+identifier+'"]').check()
                form.locator('input[name="reason"]').fill('Compose distinct equity and debt role coverage without summing capital.')
                click_and_reveal(form.get_by_role('button',name='Compose role coverage',exact=True))
                composition=page.locator('#role-composition')
                assert 'Conditional role coverage MATCH' in composition.inner_text()
                assert 'factual configuration UNRESOLVED' in composition.inner_text()
                if not args.live:
                    persisted_before=state_path.read_bytes()
                click_and_reveal(page.get_by_role('link',name='Reload',exact=True))
                if not args.live:
                    assert state_path.read_bytes() == persisted_before
                click_and_reveal(composition.get_by_role('button',name='Render Capital Alignment Brief',exact=True).first)
                with page.expect_download() as brief_download:
                    click_and_reveal(page.get_by_role('link',name='Capital Alignment Brief',exact=False))
                brief_download.value.save_as(str(output/'capital-alignment-brief.docx'))
                import docx
                brief_text='\n'.join(p.text for p in docx.Document(str(output/'capital-alignment-brief.docx')).paragraphs)
                assert 'NOT_APPROVED_BY_COMPUTATION' in brief_text and 'UNRESOLVED' in brief_text
                assert all(identifier in brief_text for identifier in role_matches)
                proof['role_composition_and_capital_brief_actual_runtime']=True
            if args.matching_games:
                from services import survey_evaluation as evaluation
                proof['matching_games']=[]
                expected={'matching:fit':'FIT', 'matching:partial':'UNRESOLVED', 'matching:mandatory-failure':'NON_FIT',
                    'matching:historical':'UNRESOLVED', 'matching:repeated-claim':'FIT',
                    'matching:missing-provenance':'UNRESOLVED', 'matching:composition':'MATCH', 'matching:brief':'UNRESOLVED'}
                for case in expected:
                    page.goto(base+'/admin/survey-evaluation')
                    page.select_option('#case',case)
                    click_and_reveal(page.get_by_role('button',name='Run isolated evaluation',exact=True))
                    page.wait_for_url('**/attention?analysis=*')
                    surface=page.locator('#role-composition' if case.endswith('composition') else '#requirement-matching')
                    phrase='Conditional role coverage ' if case.endswith('composition') else 'conditional model '
                    assert phrase+expected[case] in surface.inner_text()
                    assert 'UNRESOLVED' in surface.inner_text() and 'EVALUATION_INPUT' in surface.inner_text()
                    before=surface.inner_text()
                    click_and_reveal(page.get_by_role('link',name='Reload',exact=True))
                    assert surface.inner_text() == before
                    proof['matching_games'].append(dict(case=case, model=expected[case], factual_state='UNRESOLVED',
                        entry=urlparse(page.url).path+'?'+urlparse(page.url).query))
            page.screenshot(path=str(output/'professional-review.png'), full_page=True)
            if args.games:
                import hashlib
                from services import survey_evaluation as evaluation
                from services.case_workspace import CaseWorkspaceStore
                proof['continuum_games'] = []
                for game in evaluation.REVIEW_GAMES:
                    page.goto(base+'/admin/survey-evaluation')
                    page.select_option('#case', game)
                    click_and_reveal(page.get_by_role('button', name='Run isolated evaluation', exact=True))
                    page.wait_for_url('**/attention?analysis=*')
                    assert 'EVALUATION_INPUT' in page.locator('body').inner_text()
                    run_id = urlparse(page.url).path.split('/')[-2]
                    game_url = page.url
                    if not args.live:
                        game_path = evaluation.location(app, run_id)
                        game_record = evaluation._read(game_path)
                        game_store = CaseWorkspaceStore(game_path/'registry')
                        workspace_path = game_store._path_for(game_record['project_id'])
                        before = hashlib.sha256(workspace_path.read_bytes()).hexdigest()
                    click_and_reveal(page.get_by_role('link', name='Reload', exact=True))
                    click_and_reveal(page.locator('a').filter(has_text='Jump to evidence').first)
                    assert 'Generic' in page.locator('body').inner_text() or 'kernel' in page.locator('body').inner_text().lower()
                    page.goto(game_url)
                    if not args.live:
                        assert hashlib.sha256(workspace_path.read_bytes()).hexdigest() == before
                        workspace = game_store.get(game_record['project_id'])
                        result = workspace.analyses[-1]['governed_result']
                        assert result['evaluation_only'] and not result['canonical']
                    proof['continuum_games'].append(dict(game=game, entry_path=urlparse(game_url).path,
                        real_ui_invoked=True, source_jump=True, reload_surfaced=True,
                        persisted_bytes_unchanged=True if not args.live else None))
                page.get_by_text('Physical control paths', exact=True).scroll_into_view_if_needed()
                page.screenshot(path=str(output/'continuum-game.png'))
            if args.live:
                response=page.goto(base+'/document-shop/jobs/9c00eeec-4e65-4bde-bcea-de8b09c8beb1')
                assert response.status == 404
                proof['previously_deleted_castille_stays_absent']=True
            assert not proof['browser_errors']
            if not args.live:
                from services.runtime_observation import read
                records=[read(app, item['trace']) for item in proof['traces']]
                (output/'runtime-traces.json').write_text(json.dumps(records,indent=2),encoding='utf-8')
                owners={event['owner'] for record in records if record for event in record['events'] if event['phase']=='INVOKED'}
                for owner in ('services.image_intake.rectify_document_preview',
                    'services.document_examination.propose_text_correction',
                    'services.document_examination.review_text_correction',
                    'services.document_examination.create_working_view',
                    'services.image_intake.transform_document_preview',
                    'services.case_workspace.CaseWorkspaceStore.record_go_attention',
                    'services.case_workspace.CaseWorkspaceStore.record_temporary_relationship',
                    'services.case_workspace.CaseWorkspaceStore.run_professional_review',
                    'services.drawing_conditions.review_representation_coverage',
                    'services.case_workspace.CaseWorkspaceStore.render_professional_review',
                    'services.work_product_export.export_work_product',
                    'services.cross_modal_investigation.inspect_continuum_participation',
                    'services.quantitative_investigation.probe_interval_constraints',
                    'services.quantitative_investigation.search_interval_breakpoint',
                    'services.cross_modal_investigation.compare_normalized_information',
                    'services.quantitative_investigation.compare_scalar_values',
                    'services.document_examination.reevaluate_source_review'):
                    assert owner in owners, owner
                proof['invoked_owners']=sorted(owners)
                if args.propositions:
                    for method in ('record_review_subject', 'record_subject_proposition', 'review_subject_proposition', 'inspect_subject_propositions',
                                   'run_requirement_matching', 'inspect_requirement_matches', 'run_role_composition'):
                        assert 'services.case_workspace.CaseWorkspaceStore.'+method in owners
                    assert 'services.cross_modal_investigation.match_normalized_criteria' in owners
                    assert 'services.cross_modal_investigation.inspect_declared_temporal_scope' in owners
            (output/'proof.json').write_text(json.dumps(proof,indent=2),encoding='utf-8')
            browser.close()
    finally:
        if server: server.shutdown()
        if temporary: temporary.cleanup()
    print(json.dumps(proof,indent=2))


if __name__=='__main__': main()
