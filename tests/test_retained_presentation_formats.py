"""Every format renders the same persisted WorkProduct, never a second review."""
import copy
import hashlib
import io

import fitz
import pytest
from pptx import Presentation

from services.work_product_export import export_work_product, WorkProductExportError
from tests.test_kernel_mapping import project


def retained(project, text='Analytical coverage: PARTIAL. Governed factual status: UNRESOLVED. EVALUATION_INPUT.'):
    _, store, workspace, evidence, _ = project
    product = store.create_work_product(workspace, artifact_type='report', title='Retained review <not approval>', created_by='reviewer')
    product = store.add_work_product_section(workspace, product['id'], section_type='narrative', content={'text':text},
        content_class='human_authored', author='reviewer',
        evidence_links=[dict(object_type='evidence_item', object_id=evidence['id'])])
    return product


def extract(buffer, kind):
    if kind == 'pdf':
        with fitz.open(stream=buffer.getvalue(), filetype='pdf') as pdf:
            return '\n'.join(page.get_text() for page in pdf)
    if kind == 'pptx':
        slides = Presentation(io.BytesIO(buffer.getvalue())).slides
        return '\n'.join(shape.text for slide in slides for shape in slide.shapes if shape.has_text_frame)
    return buffer.getvalue().decode()


@pytest.mark.parametrize('kind', ['pdf', 'pptx', 'html'])
def test_formats_preserve_qualification_provenance_and_original_records(project, kind):
    _, store, workspace, evidence, _ = project
    product = retained(project)
    before = store._path_for(workspace.project_id).read_bytes()
    buffer, checksum = export_work_product(product, kind)
    text = extract(buffer, kind)
    for expected in ('PARTIAL', 'UNRESOLVED', 'EVALUATION_INPUT', 'DRAFT', evidence['id']):
        assert expected in text
    assert checksum == hashlib.sha256(buffer.getvalue()).hexdigest()
    assert store._path_for(workspace.project_id).read_bytes() == before


@pytest.mark.parametrize('kind', ['pdf', 'pptx', 'html'])
def test_long_reports_continue_without_dropping_last_unresolved_record(project, kind):
    product = retained(project, '\n'.join('Retained observation %03d: UNRESOLVED' % i for i in range(300)))
    buffer, _ = export_work_product(product, kind)
    text = extract(buffer, kind)
    for i in range(300):
        assert 'Retained observation %03d: UNRESOLVED' % i in text


@pytest.mark.parametrize('kind', ['pdf', 'pptx', 'html', 'docx', 'xlsx'])
def test_export_route_reuses_case_permissions_and_records_actual_bytes(project, kind):
    _, store, workspace, evidence, client = project
    product = retained(project)
    sources, claims = copy.deepcopy(workspace.sources), copy.deepcopy(workspace.claims)
    response = client.get(f'/projects/project/workspace/work-products/{product["id"]}/export.{kind}')
    assert response.status_code == 200
    current = store.get('project')
    saved = store.get_work_product(current, product['id'])
    assert saved['exports'][-1]['checksum'] == hashlib.sha256(response.data).hexdigest()
    assert current.sources == sources and current.claims == claims
    assert saved['sections'] == product['sections']


def test_status_banner_reports_stale_evidence_without_rewriting_prior_content(project):
    product = retained(project)
    before = copy.deepcopy(product)
    buffer, _ = export_work_product(product, 'html', status=dict(work_product={'status':'superseded'},
        evidence={'has_stale_or_broken_evidence':True}))
    assert 'SUPERSEDED' in extract(buffer, 'html') and 'REVIEW REQUIRED' in extract(buffer, 'html')
    assert product == before


def test_web_report_escapes_source_markup_and_retains_technical_fields(project):
    product = retained(project, '<script>alert("source text only")</script>')
    product['sections'][0]['content']['provenance'] = {'source_hash':'retained-hash', 'state':'UNRESOLVED'}
    buffer, _ = export_work_product(product, 'html')
    text = extract(buffer, 'html')
    assert '<script>' not in text and '&lt;script&gt;' in text
    assert '<details><summary>Supporting evidence and technical details</summary>' in text
    assert 'retained-hash' in text and 'UNRESOLVED' in text


def test_pdf_refuses_unsupported_font_instead_of_silently_losing_source_characters(project):
    product = retained(project, '未決: UNRESOLVED')
    with pytest.raises(WorkProductExportError, match='cannot preserve'):
        export_work_product(product, 'pdf')
    for kind in ('html', 'pptx'):
        buffer, _ = export_work_product(product, kind)
        assert '未決' in extract(buffer, kind)


def test_later_governing_evidence_warns_on_historical_brief_without_rewriting_it(project):
    from tests.test_reviewed_requirement_matching import matching_scope
    _, _, _, _, client = project
    store, workspace, append, run, _, _, _ = matching_scope(project)
    assert run()['state'] == 'FIT'
    product = store.render_professional_review(workspace, 'reviewer', workspace.analyses[-1]['id'])
    original = copy.deepcopy(product['sections'])
    assert not store.stale_evidence_for_work_product(workspace, product['id'])['has_stale_or_broken_evidence']
    append('new_governing_constraint', ['REVIEW'], kind='TOKEN_SET', vocabulary='review')
    before = store._path_for('project').read_bytes()
    status = store.stale_evidence_for_work_product(workspace, product['id'])
    assert status['has_stale_or_broken_evidence']
    assert any(link['object_type'] == 'analysis' for link in status['stale_links'])
    assert store._path_for('project').read_bytes() == before
    response = client.get(f'/projects/project/workspace/work-products/{product["id"]}/export.html')
    assert response.status_code == 200 and b'REVIEW REQUIRED' in response.data
    assert store.get_work_product(store.get('project'), product['id'])['sections'] == original
