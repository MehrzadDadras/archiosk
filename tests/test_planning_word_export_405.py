"""Download contract and real authenticated Flask requests; no retrieval."""
import io
import json
import re
from urllib.parse import urlsplit, parse_qs
from unittest.mock import patch

import pytest
from docx import Document
from werkzeug.security import generate_password_hash

from services import planning_studies as studies, planning_result_view as views
from services.case_workspace import CaseWorkspaceStore
from tests.test_planning_map_export import specimen


@pytest.fixture
def setup_export(tmp_path):
    import app
    from models import User, db
    from routes import planning_zoning as routes
    application = app.create_app('testing')
    application.config.update(REGISTRY_STORE_PATH=str(tmp_path), WTF_CSRF_ENABLED=False,
                              PLANNING_ZONING_LIVE_ENABLED=True)
    with application.app_context():
        db.session.add(User(username='export-planner', password_hash=generate_password_hash('test-only'), role='user'))
        db.session.commit()
    store=CaseWorkspaceStore(tmp_path); store.get_or_create('p')
    result=specimen()
    result['document']['subject']['normalized_address']='123 Synthetic Avenue'
    key=studies.WorkingResults(tmp_path).put('p','export-planner',result)
    def workspace(project):
        from flask import abort
        if project != 'p': abort(404)
        return None,store,store.get(project)
    with patch.object(routes,'_study_workspace',side_effect=workspace), patch('routes.workspace._require_export_allowed',return_value=None), patch('services.planning_live.run_live',side_effect=AssertionError('Export must not retrieve')):
        yield application,store,result,key


def sign_in(app):
    client=app.test_client()
    assert client.post('/login',data={'username':'export-planner','password':'test-only'}).status_code==302
    return client


@pytest.mark.parametrize('fmt', ['docx','pdf'])
def test_authenticated_live_download_both_methods(setup_export,fmt):
    app,store,result,key=setup_export
    client=sign_in(app)
    for method in ('get','post'):
        kwargs={'query_string' if method=='get' else 'data':dict(format=fmt,project_id='p',study_token=key)}
        response=getattr(client,method)('/planning-zoning/export',**kwargs)
        assert response.status_code==200
        assert 'attachment' in response.headers['Content-Disposition']
        if fmt=='docx':
            assert response.mimetype=='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            document=Document(io.BytesIO(response.data))
            assert result['document']['subject']['normalized_address'] in ' '.join(p.text for p in document.paragraphs)
        else:
            assert response.mimetype=='application/pdf' and response.data.startswith(b'%PDF-')
    assert not store.get('p').planning_studies  # download is not Save


def test_saved_downloads_are_exact_existing_bytes(setup_export):
    app,store,result,key=setup_export
    record=studies.WorkingResults(store.store_path).save_study(store,key,'p','export-planner',formats=['docx','pdf'])
    client=sign_in(app)
    for fmt in ('docx','pdf'):
        path=f'/planning-zoning/projects/p/studies/{key}/artifacts/report.{fmt}'
        response=client.get(path)
        assert response.status_code==200
        assert studies.maps.sha(response.data)==record['artifacts'][f'report.{fmt}']
        assert client.post(path).status_code==405
    assert len(store.get('p').planning_studies)==1


def test_login_round_trip_preserves_download_method_and_identity(setup_export):
    app,store,result,key=setup_export
    client=app.test_client()
    response=client.get('/planning-zoning/export',query_string=dict(format='docx',project_id='p',study_token=key))
    assert response.status_code==302
    login=response.headers['Location']
    next_path=parse_qs(urlsplit(login).query)['next'][0]
    assert parse_qs(urlsplit(next_path).query)['study_token']==[key]
    logged=client.post(login,data={'username':'export-planner','password':'test-only'})
    assert logged.status_code==302
    download=client.get(logged.headers['Location'])
    assert download.status_code==200
    Document(io.BytesIO(download.data))


def test_wrong_method_missing_and_expired_state_are_explicit(setup_export):
    app,store,result,key=setup_export
    client=sign_in(app)
    assert client.put('/planning-zoning/export').status_code==405
    assert client.get('/planning-zoning/export?format=docx').status_code==409
    assert client.get('/planning-zoning/export',query_string=dict(format='docx',address='1 Cassidy Pl')).status_code==409
    path=studies.WorkingResults(store.store_path)._path(key)
    envelope=json.loads(path.read_bytes());envelope['expires_at']=0
    path.write_text(json.dumps(envelope))
    assert client.get('/planning-zoning/export',query_string=dict(format='docx',project_id='p',study_token=key)).status_code==409


def test_rendered_word_and_pdf_forms_share_retained_get_contract(setup_export):
    from flask import render_template
    app,store,result,key=setup_export
    with app.test_request_context('/'):
        html=render_template('planning_zoning_result.html',view=views.build_view(result),
            exportable=True,study_token=key,study_project_id='p',panels=[],classifications=[],classification_labels={})
    for fmt in ('docx','pdf'):
        form=re.search(r'<form[^>]*data-ui-ref="planning-zoning.result.export.'+fmt+r'".*?</form>',html,re.S).group(0)
        assert 'method="get"' in form and 'action="/planning-zoning/export"' in form
        assert f'name="study_token" value="{key}"' in form
        assert 'name="project_id" value="p"' in form


def test_actual_authenticated_page_click_downloads(setup_export,tmp_path):
    from threading import Thread
    from werkzeug.serving import make_server
    from playwright.sync_api import sync_playwright
    from routes import planning_zoning as routes
    app,store,result,key=setup_export
    server=make_server('127.0.0.1',0,app,threaded=True)
    thread=Thread(target=server.serve_forever,daemon=True);thread.start()
    origin=f'http://127.0.0.1:{server.server_port}'
    live=dict(outcome='OK',view=views.build_view(result),document=result['document'],
              study_snapshot=result,timings={},source_failures=[])
    from services.planning_live import OUTCOME_OK
    live['outcome']=OUTCOME_OK
    try:
        with patch.object(routes,'_study_projects',return_value=[{'id':'p','label':'Export fixture'}]), patch('services.planning_live.run_live',return_value=live) as retrieve:
            with sync_playwright() as pw:
                browser=pw.chromium.launch(headless=True)
                page=browser.new_page(accept_downloads=True)
                page.goto(origin+'/login')
                page.locator('[name="username"]').fill('export-planner')
                page.locator('[name="password"]').fill('test-only')
                page.locator('button[type="submit"]').click()
                page.goto(origin+'/planning-zoning')
                page.locator('#planning-add-property').click()
                assert page.locator('#planning-properties [name="address"]').count()==2
                page.get_by_role('button',name='Remove property',exact=True).click()
                assert page.locator('#planning-properties [name="address"]').count()==1
                # Submit the actual analysis form once, with a retained fixture
                # standing in for the municipal response. Downloads may not read.
                form=page.locator('form').filter(has=page.locator('[name="project_id"]'))
                form.locator('[name="project_id"]').select_option('p')
                form.locator('[name="address"]').fill(result['document']['subject']['normalized_address'])
                form.locator('button[type="submit"]').click()
                page.wait_for_selector('[data-ui-ref="planning-zoning.result.export.docx-submit"]')
                assert retrieve.call_count==1
                from types import SimpleNamespace
                from services.conversational_turn import ConversationalTurnResult
                with patch('services.conversation_interpreter._evaluate_external_ai_policy', return_value=SimpleNamespace(decision='allow')), patch('services.conversational_turn.run_conversational_turn', return_value=ConversationalTurnResult(ran=True, reply_text='Scenario noted.', planning_action={'kind':'scenario'})):
                    composer=page.locator('[data-ui-ref="chat.composer"]')
                    assert composer.is_visible()
                    composer.locator('[name="text"]').fill('Consider a smaller addition')
                    composer.locator('button[type="submit"]').click()
                    page.wait_for_selector('text=Scenario updated. Municipal evidence unchanged.')
                assert retrieve.call_count==1
                for fmt in ('docx','pdf'):
                    with page.expect_download() as download_info:
                        page.locator(f'[data-ui-ref="planning-zoning.result.export.{fmt}-submit"]').click()
                    download=download_info.value
                    assert download.failure() is None
                    path=tmp_path/f'browser.{fmt}';download.save_as(path)
                    if fmt=='docx':
                        doc=Document(path)
                        assert result['document']['subject']['normalized_address'] in ' '.join(p.text for p in doc.paragraphs)
                    else: assert path.read_bytes().startswith(b'%PDF-')
                    assert retrieve.call_count==1
                browser.close()
    finally:
        server.shutdown();thread.join(timeout=5)
