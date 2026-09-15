"""Bounded Composer integration. Gateway and municipal responses are fixtures."""
import copy
import io
import json
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from docx import Document
from werkzeug.datastructures import MultiDict
from services import planning_composer as composer, planning_studies as studies
from services.conversational_turn import ConversationalTurnResult, _build_conversational_turn_prompt
from tests.test_planning_word_export_405 import setup_export, sign_in


def turn(action=None, text='See the retained source for this property.'):
    return ConversationalTurnResult(ran=True, reply_text=text, planning_action=action, grounded_in=['zone'])


def matrix(result):
    result['regulatory_matrices'] = [{'source_location': {'source_sha256': 'a'*64, 'page': 1,
        'bbox': [1,2,3,4], 'source_authority_state': 'WORKING_SOURCE'},
        'normalized_rows': [{'binding_status': 'BOUND', 'requirement_type': 'MINIMUM',
            'topic': 'Synthetic setback', 'values': {'REQUIRED': {'raw_value': '7.5 m', 'unit': 'm'},
            'EXISTING': {'raw_value': '8 m', 'unit': 'm'}, 'PROPOSED': {'raw_value': '6 m', 'unit': 'm'}}}]}]


def test_context_is_scoped_and_properties_independent(setup_export):
    _, _, result, _ = setup_export
    result['options'] = [{'option_id': 'option-A', 'summary': 'Retained option'}]
    other = copy.deepcopy(result)
    other['document']['subject']['normalized_address'] = '456 Other Avenue'
    other['document']['subject']['parcel_identifier'] = 'OTHER'
    study = composer.initialize([result, other])
    env = composer.envelope(study, 'p', 'run')
    prompt = _build_conversational_turn_prompt('Combine these properties', env, [])
    assert '123 Synthetic Avenue' in prompt and '456 Other Avenue' in prompt
    assert 'Retained option' in prompt
    assert not env.project_evidence.additional_document_evidence
    geometry = env.planning_study['properties'][0]['retrieval']['visual_geometry']['parcel']['geometry']
    assert geometry['representation'] == 'RETAINED_GEOMETRY_REFERENCE'
    assert 'coordinates' not in geometry and len(geometry['sha256']) == 64
    changed = composer.revise(study, 'Combine these properties', turn({'kind':'scenario'}), 'run')
    assert changed['document'] == result['document']
    assert changed['additional_properties'][0]['document'] == other['document']
    assert not study['workspace_context']['scenario']
    assert changed['workspace_context']['scenario'] == 'Combine these properties'


def test_numeric_refinement_uses_bound_requirement_without_authority_upgrade(setup_export):
    _, _, result, _ = setup_export
    matrix(result); study=composer.initialize([result]); before=copy.deepcopy(study)
    action={'kind':'proposal','property_index':0,'matrix_index':0,'row_index':0,'value':'4.0','unit':'m'}
    revised=composer.revise(study,'Reduce the proposal to 4.0 m',turn(action),'run')
    p=revised['workspace_context']['proposal'][0]
    assert p['required']=='7.5' and p['difference']=='3.5' and p['discrepancy']
    assert p['legal_applicability']=='UNVERIFIED' and p['status']=='PROVISIONAL'
    assert 'relief' in p['assessment'] and 'permit' in p['assessment']
    assert not p['check']['supports_established']
    assert study==before and revised['regulatory_matrices']==study['regulatory_matrices']
    for key, value in [('value','5'),('property_index',1),('unit','ft')]:
        rejected=composer.revise(study,'Reduce the proposal to 4.0 m',turn(dict(action,**{key:value})),'run')
        assert not rejected['workspace_context']['proposal']
        assert 'withheld' in rejected['workspace_context']['messages'][-1]['host_action']
    study['regulatory_matrices'][0]['normalized_rows'][0]['binding_status']='PARTIALLY_BOUND'
    assert not composer.revise(study,'4.0 m',turn(action),'run')['workspace_context']['proposal']


def test_presentation_cannot_mutate_evidence_or_hide_exception(setup_export):
    _, _, result, _ = setup_export
    study=composer.initialize([result])
    revised=composer.revise(study,'Just zoning, no map',turn({'kind':'presentation','mode':'zoning_only','show_map':False}),'run')
    assert revised['document']==study['document'] and revised['retrieval']==study['retrieval']
    report=composer.report_document(revised)
    assert not report.figures
    rows=str([t.rows for t in report.tables])
    assert 'Exception' in rows and 'text unresolved' in rows
    assert not any(t.title.endswith('Site Context') for t in report.tables)
    compact=composer.revise(study,'Make it shorter',turn({'kind':'presentation','mode':'compact'}),'run')
    full=composer.report_document(study);short=composer.report_document(compact)
    assert len(short.tables) < len(full.tables)
    assert [r for t in short.tables for r in t.rows] == [r for t in full.tables for r in t.rows]


def test_revisions_save_export_reopen_without_retrieval(setup_export):
    app,store,result,_=setup_export;client=sign_in(app)
    study=composer.initialize([result]);working=studies.WorkingResults(store.store_path)
    key=working.put('p','export-planner',study)
    base=f'/planning-zoning/projects/p/working/{key}'
    response=client.get(base)
    assert response.status_code==200 and b'Tell GO what you are considering' in response.data
    assert b'id="dock-composer-pen-sheet"' not in response.data
    assert b'id="dock-composer-image"' not in response.data
    assert b'workspace.start_new_conversation' not in response.data
    with patch('services.conversation_interpreter._evaluate_external_ai_policy',return_value=SimpleNamespace(decision='allow')), patch('services.conversational_turn.run_conversational_turn',return_value=turn({'kind':'scenario'})) as go:
        response=client.post(base+'/conversation',data={'text':'Consider a smaller addition'})
    assert response.status_code==302
    newer=response.headers['Location'].rsplit('/',1)[1]
    assert newer!=key and not working.get(key,'p','export-planner')['workspace_context']['messages']
    assert go.call_args.kwargs['envelope'].planning_study['project_id']=='p'
    for fmt in ('pdf','docx'):
        exported=client.get('/planning-zoning/export',query_string={'project_id':'p','study_token':newer,'format':fmt})
        assert exported.status_code==200
        if fmt=='docx':
            doc=Document(io.BytesIO(exported.data))
            assert 'Consider a smaller addition' in str([[c.text for c in row.cells] for t in doc.tables for row in t.rows])
        else: assert exported.data.startswith(b'%PDF')
    response=client.post('/planning-zoning/projects/p/studies',data=MultiDict([('study_token',newer),('formats','pdf'),('formats','docx')]))
    assert response.status_code==302
    record=store.get('p').planning_studies[0]
    frozen=studies.artifact(store,store.get('p'),newer,'snapshot.json')
    assert client.get(response.headers['Location']).status_code==200
    assert frozen==studies.artifact(store,store.get('p'),newer,'snapshot.json')
    assert json.loads(frozen)['workspace_context']['scenario']=='Consider a smaller addition'
    assert client.post(f'/planning-zoning/projects/p/studies/{newer}/continue').status_code==302
    assert record['snapshot_sha256']==store.get('p').planning_studies[0]['snapshot_sha256']


def test_cross_project_actor_and_policy_block_before_go(setup_export):
    app,store,result,_=setup_export;client=sign_in(app)
    working=studies.WorkingResults(store.store_path)
    for project,actor in [('other','export-planner'),('p','other-user')]:
        key=working.put(project,actor,composer.initialize([result]))
        with patch('services.conversational_turn.run_conversational_turn') as go:
            assert client.post(f'/planning-zoning/projects/p/working/{key}/conversation',data={'text':'Explain'}).status_code==409
            go.assert_not_called()
    key=working.put('p','export-planner',composer.initialize([result]))
    with patch('services.conversation_interpreter._evaluate_external_ai_policy',return_value=SimpleNamespace(decision='DENY')), patch('services.conversational_turn.run_conversational_turn') as go:
        assert client.post(f'/planning-zoning/projects/p/working/{key}/conversation',data={'text':'Explain'}).status_code==403
        go.assert_not_called()


def test_explanation_uses_existing_turn_without_mutation(setup_export):
    _,_,result,_=setup_export
    study=composer.initialize([result])
    revised=composer.revise(study,'Where did this zone come from?',turn(),'run')
    assert revised['document']==study['document']
    assert revised['workspace_context']['messages'][-1]['grounded_in']==['zone']
    assert not revised['workspace_context']['scenario']


def test_multi_intake_and_investigation_use_normal_route(setup_export):
    from services.planning_live import OUTCOME_OK
    app,store,result,_=setup_export;client=sign_in(app)
    second=copy.deepcopy(result);second['document']['subject']['normalized_address']='456 Second Street'
    def live(address):
        return {'outcome':OUTCOME_OK,'study_snapshot':copy.deepcopy(result if address.startswith('123') else second)}
    with patch('services.planning_live.run_live',side_effect=live) as retrieve:
        response=client.post('/planning-zoning/analyze',data=MultiDict([('composer','1'),('project_id','p'),('address','123 Synthetic Avenue'),('address','456 Second Street')]))
        assert response.status_code==302 and retrieve.call_count==2
    key=response.headers['Location'].rsplit('/',1)[1]
    working=studies.WorkingResults(store.store_path);study=working.get(key,'p','export-planner')
    assert len(composer.properties(study))==2
    revised=composer.revise(study,'Find the exception text for property two',turn({'kind':'investigate','property_index':1}),key)
    pending=working.put('p','export-planner',revised)
    with patch('services.planning_live.run_live',side_effect=live) as retrieve:
        response=client.post(f'/planning-zoning/projects/p/working/{pending}/investigate')
        assert response.status_code==302
        retrieve.assert_called_once_with('456 Second Street')
    assert working.get(key,'p','export-planner')==study
    record=working.save_study(store,key,'p','export-planner',formats=['docx'])
    assert 'evidence/property-2/zoning-map.png' in record['artifacts']
    assert len(record['properties'])==2


def test_generic_composer_unchanged_and_planning_action_is_opt_in(setup_export):
    from services.conversational_turn import ContextEnvelope, ProjectEvidence, run_conversational_turn
    _,store,result,_=setup_export
    ordinary=ContextEnvelope(None,None,None,None,ProjectEvidence('ordinary.txt'))
    assert 'planning_action' not in _build_conversational_turn_prompt('Explain',ordinary,[])
    outcome=SimpleNamespace(ran=True,parsed={'reply_text':'Source says R','planning_action':{'kind':'scenario'}},provider='fixture',model='fixture',requested_at='now')
    with patch('services.conversational_turn.call_llm_json',return_value=outcome):
        assert run_conversational_turn('Explain',store.get('p'),ordinary).planning_action is None
        scoped=composer.envelope(composer.initialize([result]),'p','run')
        assert run_conversational_turn('Explain',store.get('p'),scoped).planning_action=={'kind':'scenario'}


def test_presentation_cannot_hide_visual_disagreement(setup_export):
    _,_,result,_=setup_export
    study=composer.initialize([result])
    study['document']['statements'][0]['text']='The parcel lies within WRONG.'
    study['workspace_context']['presentation']['show_map']=False
    with pytest.raises(ValueError,match='disagreement'):
        composer.report_document(study)


def test_conclusion_and_exception_detail_survive_projection(setup_export):
    from services import planning_report, planning_result_view
    _,_,result,_=setup_export
    result['conclusion']='Retained bounded conclusion.'
    result['document']['site_specific_exceptions']=[{'exception_id':'8','text_retrieved':True,
        'development_effect':'Keep the established exception condition.'}]
    before=copy.deepcopy(result)
    projection=planning_report.build(planning_result_view.build_view(result),result)
    rows=str(projection['sections'])
    assert result['conclusion'] in rows and 'Keep the established exception condition.' in rows
    assert result==before
